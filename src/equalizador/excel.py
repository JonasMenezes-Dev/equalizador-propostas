from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _preparar_mapa_compacto(mapa: pd.DataFrame) -> pd.DataFrame:
    """Cria uma leitura vertical: uma linha por item e fornecedor."""
    colunas_base = {
        "ID", "Descrição Solicitada", "Quantidade", "Unidade",
        "Menor Preço", "Maior Preço", "Diferença %",
        "Melhor Fornecedor", "Status Equalização",
    }
    fornecedores = [coluna for coluna in mapa.columns if coluna not in colunas_base]
    linhas = []

    for _, item in mapa.iterrows():
        status_por_fornecedor = {}
        for status in str(item.get("Status Equalização", "")).split(" | "):
            if ": " in status:
                fornecedor, valor = status.split(": ", 1)
                status_por_fornecedor[fornecedor] = valor

        for fornecedor in fornecedores:
            preco = item.get(fornecedor)
            status = status_por_fornecedor.get(fornecedor, "NÃO COTADO")
            observacao = "Melhor preço"
            if fornecedor not in str(item.get("Melhor Fornecedor", "")).split(", "):
                observacao = ""
            linhas.append({
                "ID": item.get("ID"),
                "Descrição": item.get("Descrição Solicitada"),
                "Quantidade": item.get("Quantidade"),
                "Unidade": item.get("Unidade"),
                "Fornecedor": fornecedor,
                "Preço Unitário": preco,
                "Status": status,
                "Melhor Fornecedor": item.get("Melhor Fornecedor"),
                "Observação": observacao,
            })

    return pd.DataFrame(linhas, columns=[
        "ID", "Descrição", "Quantidade", "Unidade", "Fornecedor",
        "Preço Unitário", "Status", "Melhor Fornecedor", "Observação",
    ])


def _preparar_precos_verticais(mapa: pd.DataFrame) -> pd.DataFrame:
    colunas_base = {
        "ID", "Descrição Solicitada", "Quantidade", "Unidade",
        "Menor Preço", "Maior Preço", "Diferença %",
        "Melhor Fornecedor", "Status Equalização",
    }
    fornecedores = [coluna for coluna in mapa.columns if coluna not in colunas_base]
    if not fornecedores:
        return pd.DataFrame(columns=["ID", "Descrição", "Fornecedor", "Preço"])

    precos = mapa.melt(
        id_vars=["ID", "Descrição Solicitada"],
        value_vars=fornecedores,
        var_name="Fornecedor",
        value_name="Preço",
    ).dropna(subset=["Preço"])
    return precos.rename(columns={"Descrição Solicitada": "Descrição"})[
        ["ID", "Descrição", "Fornecedor", "Preço"]
    ].sort_values(["ID", "Preço", "Fornecedor"], kind="mergesort")


def gerar_relatorio_excel(
    caminho: str | Path,
    *,
    indicadores: dict,
    insights: list[str],
    mapa_comparativo: pd.DataFrame,
    resumo_fornecedores: pd.DataFrame,
    alertas: pd.DataFrame,
    indicadores_comerciais: pd.DataFrame,
    recomendacoes: pd.DataFrame,
    base_recomendacao: pd.DataFrame,
    analise_executiva: dict,
    analise_fonte: dict | None = None,
    descartes: list[dict] | None = None,
) -> Path:
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)

    indicadores_df = pd.DataFrame(
        [{"Indicador": k, "Valor": v}
         for k, v in indicadores.items()]
    )
    insights_df = pd.DataFrame({"Insight": insights})
    executivo_df = pd.DataFrame(
        [
            {"Indicador": chave, "Valor": valor}
            for chave, valor in analise_executiva.items()
            if not isinstance(valor, pd.DataFrame)
        ]
    )
    ranking_executivo = analise_executiva.get("Ranking por Valor", pd.DataFrame())
    mapa_compacto = _preparar_mapa_compacto(mapa_comparativo)
    precos_verticais = _preparar_precos_verticais(mapa_comparativo)

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        executivo_df.to_excel(writer, sheet_name="Dashboard", index=False)
        indicadores_df.to_excel(writer, sheet_name="Resumo", index=False)
        executivo_df.to_excel(writer, sheet_name="Análise Executiva", index=False)
        if isinstance(ranking_executivo, pd.DataFrame):
            ranking_executivo.to_excel(writer, sheet_name="Ranking Executivo", index=False)
        insights_df.to_excel(writer, sheet_name="Insights", index=False)
        mapa_compacto.to_excel(writer, sheet_name="Mapa Comparativo", index=False)
        precos_verticais.to_excel(writer, sheet_name="Preços por Fornecedor", index=False)
        resumo_fornecedores.to_excel(writer, sheet_name="Fornecedores", index=False)
        alertas.to_excel(writer, sheet_name="Alertas", index=False)
        indicadores_comerciais.to_excel(
            writer, sheet_name="Condições Comerciais", index=False
        )
        recomendacoes.to_excel(writer, sheet_name="Recomendações", index=False)
        base_recomendacao.to_excel(writer, sheet_name="Base Recomendação", index=False)
        descartes_df = pd.DataFrame(
            descartes or [],
            columns=["Linha (Excel)", "Motivo", "Descrição encontrada", "Fornecedor encontrado"],
        )
        descartes_df.to_excel(writer, sheet_name="Linhas Descartadas", index=False)
        if analise_fonte:
            perfil = analise_fonte.get("perfil", {})
            pd.DataFrame(
                [{"Indicador": key, "Valor": value} for key, value in perfil.get("qualidade", {}).items()]
                + [{"Indicador": "Tipo de fonte", "Valor": perfil.get("tipo")}]
                + [{"Indicador": "Observação", "Valor": perfil.get("observacao")}]
            ).to_excel(writer, sheet_name="Perfil da Fonte", index=False)
            analise_fonte.get("produtos", pd.DataFrame()).to_excel(
                writer, sheet_name="Produtos Analisados", index=False
            )
            analise_fonte.get("fornecedores", pd.DataFrame()).to_excel(
                writer, sheet_name="Fornecedores Analisados", index=False
            )
            matriz = analise_fonte.get("matriz_referencia", {})
            if matriz:
                matriz.get("tabela", pd.DataFrame()).to_excel(
                    writer, sheet_name="Referência da Matriz", index=False
                )
                pd.DataFrame([
                    {"Indicador": "Tipo", "Valor": matriz.get("tipo")},
                    {"Indicador": "Fornecedores de propostas iniciais", "Valor": ", ".join(matriz.get("fornecedores", []))},
                ]).to_excel(writer, sheet_name="Metadados da Matriz", index=False)

    formatar_relatorio(caminho)
    return caminho


def formatar_relatorio(caminho: str | Path) -> None:
    wb = load_workbook(caminho)

    header_fill = PatternFill("solid", fgColor="173F5F")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    stripe_fill = PatternFill("solid", fgColor="F5F8FA")
    border = Border(bottom=Side(style="hair", color="D9E2E8"))
    fills = {
        "bad": PatternFill("solid", fgColor="F4CCCC"),
        "medium": PatternFill("solid", fgColor="FFF2CC"),
        "good": PatternFill("solid", fgColor="D9EAD3"),
    }
    tab_colors = {
        "Resumo": "2A9D8F",
        "Análise Executiva": "2A9D8F",
        "Ranking Executivo": "2A9D8F",
        "Recomendações": "E9C46A",
        "Alertas": "E76F51",
        "Linhas Descartadas": "C97B63",
    }

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = tab_colors.get(ws.title, "8BA6B5")
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        ws.row_dimensions[1].height = 24

        for column_index, col in enumerate(ws.iter_cols(), start=1):
            max_len = max(
                (len(str(c.value)) for c in col if c.value is not None),
                default=0,
            )
            header = str(ws.cell(1, column_index).value or "")
            minimum = 12 if header not in {
                "Descrição", "Descrição Solicitada", "Descrição Cotada",
                "Motivo", "Insight", "Mensagem", "Observação",
            } else 24
            width = min(max(max_len + 2, minimum), 36)
            ws.column_dimensions[get_column_letter(column_index)].width = width

        for row in ws.iter_rows(min_row=2):
            row_number = row[0].row
            if row_number is None:
                continue
            if ws.title not in {"Resumo", "Análise Executiva"} and row_number % 2 == 0:
                for cell in row:
                    cell.fill = stripe_fill
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 3})
            ws.row_dimensions[row_number].height = 30 if ws.title in {"Insights", "Alertas"} else 20

    moeda_headers = {
        "Preço", "Preço Unitário", "Valor Unitário", "Menor Preço", "Maior Preço",
        "Preço Mínimo", "Preço Médio", "Preço Máximo",
        "Total Proposta Elegível", "Economia Potencial",
        "Total Comprando pelo Menor Preço",
    }

    for ws in wb.worksheets:
        for column_index, cell in enumerate(ws[1], start=1):
            if cell.value in moeda_headers:
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=column_index).number_format = 'R$ #,##0.00'

        if ws.title == "Dashboard":
            dashboard_moeda = {"Economia Potencial", "Total Comprando pelo Menor Preço"}
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=1).value in dashboard_moeda:
                    ws.cell(row=row, column=2).number_format = 'R$ #,##0.00'

        for column_index, cell in enumerate(ws[1], start=1):
            if "%" in str(cell.value):
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=column_index).number_format = '0.00"%"'

        _colorir_indicadores(ws, fills, stripe_fill)

    wb.save(caminho)


def _colorir_indicadores(
    ws, fills: dict[str, PatternFill], stripe_fill: PatternFill
) -> None:
    """Aplica cores semânticas somente às colunas de avaliação do relatório."""
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(ws[1], start=1)
        if cell.value is not None
    }

    for header, column_index in headers.items():
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=column_index)
            value = str(cell.value).strip().upper() if cell.value is not None else ""

            if header in {"Status", "Status Equalização", "Validação Quantidade", "Validação Unidade"}:
                if value in {"NÃO COTADO", "DIVERGENTE", "DIVERGÊNCIA", "DIVERGÊNCIA ESTRUTURAL"}:
                    cell.fill = fills["bad"]
                elif value in {"EQUIVALENTE COM REVISÃO", "REVISÃO", "SOB CONSULTA"}:
                    cell.fill = fills["medium"]
                elif value in {"OK", "EQUIVALENTE"}:
                    cell.fill = fills["good"]

            elif header in {"Classificação Recomendação", "Classificação"}:
                if "NÃO" in value or "SEM OPÇÃO" in value:
                    cell.fill = fills["bad"]
                elif value in {"ACEITÁVEL", "BOA"}:
                    cell.fill = fills["medium"]
                elif value in {"EXCELENTE"}:
                    cell.fill = fills["good"]

            elif header == "Elegível para Recomendação":
                if value in {"FALSE", "NÃO", "0"}:
                    cell.fill = fills["bad"]
                elif value in {"TRUE", "SIM", "1"}:
                    cell.fill = fills["good"]

            elif header in {
                "Score Preço", "Score Similaridade", "Score Atributos",
                "Score Comercial", "Score Recomendação",
            }:
                if isinstance(cell.value, (int, float)):
                    cell.fill = fills["good"] if cell.value >= 90 else (
                        fills["medium"] if cell.value >= 60 else stripe_fill
                    )

            elif header == "Diferença %" and isinstance(cell.value, (int, float)):
                cell.fill = (
                    fills["good"] if cell.value <= 5
                    else fills["medium"] if cell.value <= 15
                    else fills["bad"]
                )
