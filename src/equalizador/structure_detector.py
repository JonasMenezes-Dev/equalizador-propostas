from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
import re

import pandas as pd
from openpyxl import load_workbook

from .diagnostics import DiagnosticoPlanilha, DICAS_PADRAO
from .semantic_engine import SemanticEngine


@dataclass(frozen=True)
class ColumnDetection:
    position: int
    original_name: str
    concept: str | None
    score: float
    confidence: str
    matched_alias: str | None
    data_score: float = 0.0
    evidence: str = ""


@dataclass
class SheetDetection:
    sheet_name: str
    rows: int
    columns: int
    header_row: int
    data_start_row: int
    header_score: float
    header_values: list[str] = field(default_factory=list)
    columns_detected: list[ColumnDetection] = field(default_factory=list)
    supplier_candidates: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def concept_map(self) -> dict[str, list[str]]:
        result = {}
        for item in self.columns_detected:
            if item.concept:
                result.setdefault(item.concept, []).append(item.original_name)
        return result


@dataclass
class WorkbookDetection:
    file_path: str
    selected_sheet: str | None
    sheets: list[SheetDetection]
    warnings: list[str] = field(default_factory=list)

    @property
    def selected(self) -> SheetDetection | None:
        return next(
            (s for s in self.sheets if s.sheet_name == self.selected_sheet),
            None,
        )


class StructureDetector:
    REQUIRED_CONCEPTS = {"DESCRIPTION", "UNIT_PRICE"}
    USEFUL_CONCEPTS = {
        "ID", "DESCRIPTION", "QUANTITY", "UNIT", "UNIT_PRICE",
        "TOTAL_PRICE", "SUPPLIER", "PAYMENT_DAYS", "DELIVERY_DAYS",
        "FREIGHT", "MIN_ORDER", "PROPOSAL_VALIDITY", "BENCHMARK",
        "NEW_RATE", "DIFFERENCE", "REDUCTION", "INCREASE", "STATUS",
        "OBSERVATION",
    }

    def __init__(
        self,
        semantic_engine: SemanticEngine | None = None,
        max_header_rows: int = 30,
        sample_rows: int = 100,
    ):
        self.semantic = semantic_engine or SemanticEngine()
        self.max_header_rows = max_header_rows
        self.sample_rows = sample_rows

    @staticmethod
    def _clean_header(value: Any, position: int) -> str:
        if value is None or pd.isna(value):
            return f"Unnamed_{position + 1}"
        text = str(value).strip()
        return text if text else f"Unnamed_{position + 1}"

    @staticmethod
    def _empty(value: Any) -> bool:
        if value is None:
            return True
        try:
            return bool(pd.isna(value))
        except (TypeError, ValueError):
            return False

    @classmethod
    def _data_profile(cls, values: pd.Series) -> dict[str, float]:
        values = values[~values.map(cls._empty)]
        if values.empty:
            return {
                "filled": 0.0, "numeric": 0.0, "integer": 0.0,
                "text": 0.0, "unique_ratio": 0.0, "median_length": 0.0,
                "currency": 0.0,
            }

        text = values.map(lambda value: str(value).strip())
        def parse_number(value: str) -> float | None:
            if not re.fullmatch(r"[\sR$€£+-]*\d[\d.,\s]*", value, re.I):
                return None
            normalized = re.sub(r"[^0-9,.-]", "", value)
            comma, dot = normalized.rfind(","), normalized.rfind(".")
            if comma >= 0 and dot >= 0:
                if comma > dot:
                    normalized = normalized.replace(".", "").replace(",", ".")
                else:
                    normalized = normalized.replace(",", "")
            elif comma >= 0:
                normalized = normalized.replace(",", ".")
            try:
                return float(normalized)
            except ValueError:
                return None

        numbers = text.map(parse_number)
        numeric_mask = numbers.notna()
        return {
            "filled": 1.0,
            "numeric": float(numeric_mask.mean()),
            "integer": float((numbers[numeric_mask] % 1 == 0).mean()) if numeric_mask.any() else 0.0,
            "text": float((~numeric_mask).mean()),
            "unique_ratio": float(text.nunique(dropna=True) / len(text)),
            "median_length": float(text.str.len().median()),
            "currency": float(text.str.contains(r"R\$|\$|€|£", case=False, regex=True).mean()),
        }

    @classmethod
    def _data_evidence(cls, concept: str | None, values: pd.Series) -> tuple[float, str]:
        profile = cls._data_profile(values)
        if not profile["filled"]:
            return 0.0, "sem dados"

        unit_values = {
            "un", "und", "unid", "unid.", "pc", "pct", "kg", "g", "l",
            "ml", "m", "cx", "caixa", "pacote", "pç", "hora", "h",
        }
        normalized_values = {
            str(value).strip().casefold().rstrip(".")
            for value in values if not cls._empty(value)
        }
        unit_ratio = len(normalized_values & unit_values) / max(len(normalized_values), 1)

        if concept == "UNIT":
            score = 24 * unit_ratio + (4 if profile["unique_ratio"] <= 0.2 else 0)
            return min(score, 26.0), "valores parecem unidades"
        if concept == "DESCRIPTION":
            score = 12 * profile["text"] + (6 if profile["median_length"] >= 8 else 0)
            score += 4 if profile["unique_ratio"] >= 0.75 else 0
            return min(score, 22.0), "texto descritivo"
        if concept == "SUPPLIER":
            repeated = profile["unique_ratio"] < 0.8 and len(normalized_values) >= 2
            score = 12 * profile["text"] + (8 if repeated else 0)
            return min(score, 22.0), "valores textuais repetidos"
        if concept == "ID":
            score = 12 * profile["text"] + (8 if profile["unique_ratio"] >= 0.8 else 0)
            score += 2 if profile["median_length"] <= 40 else 0
            return min(score, 22.0), "valores com aparência de identificador"
        if concept in {"QUANTITY", "UNIT_PRICE", "TOTAL_PRICE", "PAYMENT_DAYS", "DELIVERY_DAYS"}:
            score = 16 * profile["numeric"]
            if concept in {"UNIT_PRICE", "TOTAL_PRICE"}:
                score += 5 * max(profile["currency"], 0.0)
                score += 12 * (1 - profile["integer"])
            elif concept == "QUANTITY":
                score += 4 * profile["integer"]
            return min(score, 22.0), "valores numéricos"
        return 0.0, ""

    @classmethod
    def _infer_concept(cls, header: str, values: pd.Series) -> tuple[str | None, float, str]:
        candidates = [
            "DESCRIPTION", "SUPPLIER", "UNIT", "ID", "QUANTITY", "UNIT_PRICE",
        ]
        scored = [
            (concept, *cls._data_evidence(concept, values))
            for concept in candidates
        ]
        concept, score, evidence = max(scored, key=lambda item: item[1])
        if score < 14:
            return None, 0.0, ""
        return concept, score, f"cabeçalho genérico; {evidence}"

    def _candidate_score(self, row: pd.Series, total_columns: int) -> float:
        values = [value for value in row.tolist() if not self._empty(value)]
        headers = [
            self._clean_header(v, i)
            for i, v in enumerate(row.tolist())
        ]
        headers = [h for h in headers if not h.startswith("Unnamed_")]
        if not headers:
            return 0.0

        matches = [self.semantic.classify(h) for h in headers]
        concepts = {m.concept for m in matches if m.concept}
        high = sum(m.confidence == "HIGH" for m in matches)
        medium = sum(m.confidence == "MEDIUM" for m in matches)
        density = len(headers) / max(total_columns, 1)

        score = high * 18 + medium * 8 + len(concepts) * 5 + density * 10

        if self.REQUIRED_CONCEPTS.issubset(concepts):
            score += 35
        elif concepts & self.REQUIRED_CONCEPTS:
            score += 12

        numeric_values = pd.to_numeric(
            pd.Series(values).astype(str).str.replace(r"[^0-9,.-]", "", regex=True).str.replace(",", ".", regex=False),
            errors="coerce",
        )
        numeric_ratio = float(numeric_values.notna().mean()) if values else 0.0
        if numeric_ratio >= 0.25:
            score -= 55

        return score

    def _find_header_row(self, df: pd.DataFrame) -> tuple[int, float]:
        best_row, best_score = 0, 0.0
        for idx in range(min(self.max_header_rows, len(df))):
            headers = [
                self._clean_header(value, position)
                for position, value in enumerate(df.iloc[idx].tolist())
            ]
            generic_count = sum(
                bool(re.fullmatch(r"campo\s+[a-z]+", header, re.I))
                for header in headers
            )
            if generic_count >= 6:
                return idx, 40.0

            score = self._candidate_score(df.iloc[idx], df.shape[1])
            adjusted = score - idx * 0.15
            if adjusted > best_score:
                best_row, best_score = idx, adjusted

        return best_row, round(best_score, 2)

    def _detect_columns(self, df: pd.DataFrame, header_row: int):
        headers = [
            self._clean_header(v, i)
            for i, v in enumerate(df.iloc[header_row].tolist())
        ]
        result = []
        values = df.iloc[header_row + 1:]
        context = headers
        for i, header in enumerate(headers):
            m = self.semantic.classify(header, context=context)
            data_score, evidence = self._data_evidence(m.concept, values.iloc[:, i])
            concept = m.concept
            if m.confidence in {"LOW", "NONE"}:
                inferred_concept, inferred_score, inferred_evidence = self._infer_concept(
                    header, values.iloc[:, i]
                )
                if inferred_concept is not None:
                    concept = inferred_concept
                    data_score = inferred_score
                    evidence = inferred_evidence
            result.append(ColumnDetection(
                i, header, concept, m.score, m.confidence, m.matched_alias,
                round(data_score, 2), evidence,
            ))
        return result

    def _find_data_start(self, df: pd.DataFrame, header_row: int) -> int:
        for idx in range(header_row + 1, len(df)):
            non_empty = sum(not self._empty(v) for v in df.iloc[idx].tolist())
            if non_empty >= 2:
                return idx + 1
        return min(header_row + 2, len(df) + 1)

    @staticmethod
    def _suppliers(columns: list[ColumnDetection]) -> list[str]:
        found=set()
        pattern=re.compile(
            r"^(?P<supplier>.+?)\s*(?:-|–|—|\||:)\s*"
            r"(?:preco|preço|valor|price|unit price|rate|quantidade|qtd|"
            r"quantity|unidade|unit|descricao|descrição|description)\b",
            re.I,
        )
        for c in columns:
            m=pattern.search(c.original_name.strip())
            if m:
                name=m.group("supplier").strip(" -–—|_:")
                if name:
                    found.add(name)
        return sorted(found)

    def inspect_sheet(self, file_path: str | Path, sheet_name: str) -> SheetDetection:
        df = pd.read_excel(
            file_path, sheet_name=sheet_name, header=None,
            nrows=self.sample_rows, engine="openpyxl", dtype=object
        )
        if df.empty:
            return SheetDetection(sheet_name,0,0,1,1,0.0,warnings=["Aba vazia."])

        header_idx, score = self._find_header_row(df)
        columns = self._detect_columns(df, header_idx)
        data_start = self._find_data_start(df, header_idx)

        concepts={c.concept for c in columns if c.concept}
        warnings=[]
        missing=self.REQUIRED_CONCEPTS-concepts
        if missing:
            warnings.append("Conceitos centrais não identificados: "+", ".join(sorted(missing)))
        if score < 35:
            warnings.append("Confiança baixa na identificação do cabeçalho.")
        if score <= 0:
            # Nenhum cabeçalho pontuou: o motor caiu na linha 0 por falta de
            # evidência. Isso indica que a planilha pode não ter cabeçalho ou
            # estar em um layout atípico — o usuário precisa ser avisado.
            warnings.append(
                "Nenhum cabeçalho com sinais reconhecíveis foi encontrado; "
                f"foi assumida a linha {header_idx + 1} como cabeçalho. "
                "Confira se a aba correta foi escolhida ou informe o nome da aba."
            )
        inferred = [column.original_name for column in columns if column.evidence.startswith("cabeçalho genérico")]
        if inferred:
            warnings.append(
                "Campos inferidos pelo conteúdo das células: " + ", ".join(inferred) + "."
            )

        return SheetDetection(
            sheet_name=sheet_name,
            rows=len(df),
            columns=df.shape[1],
            header_row=header_idx+1,
            data_start_row=data_start,
            header_score=score,
            header_values=[c.original_name for c in columns],
            columns_detected=columns,
            supplier_candidates=self._suppliers(columns),
            warnings=warnings,
        )

    def inspect(self, file_path: str | Path) -> WorkbookDetection:
        path=Path(file_path)
        if not path.exists():
            raise DiagnosticoPlanilha(
                mensagem=f"O arquivo não foi encontrado em: {path}",
                codigo="arquivo_ausente",
                titulo="Arquivo indisponível",
                dicas=["Envie novamente a planilha."],
            )
        if path.suffix.lower() not in {".xlsx",".xlsm",".xltx",".xltm"}:
            raise DiagnosticoPlanilha(
                mensagem=(
                    f"O formato '{path.suffix or 'sem extensão'}' não é suportado. "
                    "Use arquivos .xlsx, .xlsm, .xltx ou .xltm."
                ),
                codigo="formato_nao_suportado",
                titulo="Formato de arquivo não suportado",
                dicas=[
                    "Abra o arquivo no Excel e salve como .xlsx.",
                    "Formatos antigos como .xls e .csv não são aceitos.",
                ],
            )

        try:
            wb=load_workbook(path, read_only=True, data_only=True)
            names=wb.sheetnames
            wb.close()
        except DiagnosticoPlanilha:
            raise
        except Exception as exc:
            raise DiagnosticoPlanilha(
                mensagem=(
                    "O arquivo não pôde ser aberto como planilha Excel. "
                    f"Ele pode estar corrompido ou protegido: {exc}"
                ),
                codigo="arquivo_corrompido",
                titulo="Arquivo ilegível",
                dicas=[
                    "Abra o arquivo no Excel e salve uma nova cópia em .xlsx.",
                    "Remova a proteção por senha antes de enviar.",
                ],
            ) from exc

        if not names:
            raise DiagnosticoPlanilha(
                mensagem="O arquivo não possui nenhuma aba.",
                codigo="sem_abas",
                titulo="Arquivo sem abas",
                dicas=["Confirme que a planilha contém dados em pelo menos uma aba."],
            )

        sheets=[]
        warnings=[]
        for name in names:
            try:
                sheets.append(self.inspect_sheet(path,name))
            except Exception as exc:
                warnings.append(f"Aba '{name}' não pôde ser analisada: {exc}")

        if not sheets:
            raise DiagnosticoPlanilha(
                mensagem=(
                    "Nenhuma das abas pôde ser analisada. "
                    f"Abas do arquivo: {', '.join(names)}."
                ),
                codigo="nenhuma_aba_analisavel",
                titulo="Nenhuma aba analisável",
                abas_analisadas=list(names),
                dicas=list(warnings) + DICAS_PADRAO,
            )

        def rank(s):
            concepts=set(s.concept_map)
            return (
                len(concepts & self.REQUIRED_CONCEPTS)*1000
                + len(concepts & self.USEFUL_CONCEPTS)*50
                + sum(c.confidence=="HIGH" for c in s.columns_detected)*5
                + sum(c.data_score for c in s.columns_detected)
                + s.header_score,
                s.rows,
                s.columns,
            )

        selected=max(sheets,key=rank)
        if not self.REQUIRED_CONCEPTS.issubset(set(selected.concept_map)):
            warnings.append(
                f"Aba '{selected.sheet_name}' selecionada com baixa confiança "
                "para DESCRIPTION + UNIT_PRICE."
            )

        return WorkbookDetection(str(path),selected.sheet_name,sheets,warnings)

    def summary(self, result: WorkbookDetection) -> str:
        s=result.selected
        if s is None:
            return "Nenhuma aba selecionada."
        lines=[
            f"Arquivo: {result.file_path}",
            f"Aba selecionada: {s.sheet_name}",
            f"Cabeçalho: linha {s.header_row}",
            f"Início dos dados: linha {s.data_start_row}",
            f"Score estrutural: {s.header_score:.2f}",
            f"Fornecedores: {', '.join(s.supplier_candidates) or 'não identificados'}",
            "Conceitos detectados:",
        ]
        for concept, cols in sorted(s.concept_map.items()):
            lines.append(f"  - {concept}: {', '.join(cols)}")
        for warning in s.warnings + result.warnings:
            lines.append(f"  ⚠ {warning}")
        return "\n".join(lines)
